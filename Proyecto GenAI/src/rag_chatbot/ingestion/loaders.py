from datetime import datetime, timezone
from pathlib import Path

import fitz
from docx import Document

from rag_chatbot.config import PROJECT_ROOT
from rag_chatbot.schemas import DocumentPage, DocumentStatus, ProcessedDocument


MAX_ERROR_MESSAGE_LENGTH = 500
TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")


def load_document(path: Path, file_hash: str, document_id: str) -> ProcessedDocument:
    file_type = path.suffix.lower().lstrip(".")

    try:
        if file_type == "pdf":
            return _load_pdf(path, file_hash, document_id)
        if file_type == "txt":
            return _load_txt(path, file_hash, document_id)
        if file_type == "docx":
            return _load_docx(path, file_hash, document_id)

        return _build_document(
            path=path,
            file_hash=file_hash,
            document_id=document_id,
            status=DocumentStatus.UNSUPPORTED,
            error_message=f"Unsupported file type: {file_type}",
        )
    except Exception as exc:
        return _build_document(
            path=path,
            file_hash=file_hash,
            document_id=document_id,
            status=DocumentStatus.FAILED,
            error_message=_safe_error(exc),
        )


def _load_pdf(path: Path, file_hash: str, document_id: str) -> ProcessedDocument:
    pages: list[DocumentPage] = []

    with fitz.open(path) as document:
        for page_index, page in enumerate(document, start=1):
            text = _clean_text(page.get_text("text") or "")
            pages.append(
                DocumentPage(
                    page_number=page_index,
                    text=text,
                    char_count=len(text),
                )
            )

    char_count = sum(page.char_count for page in pages)
    requires_ocr = len(pages) > 0 and char_count == 0
    status = DocumentStatus.EMPTY_TEXT if char_count == 0 else DocumentStatus.EXTRACTED

    return _build_document(
        path=path,
        file_hash=file_hash,
        document_id=document_id,
        status=status,
        requires_ocr=requires_ocr,
        pages=pages,
        page_count=len(pages),
        char_count=char_count,
    )


def _load_txt(path: Path, file_hash: str, document_id: str) -> ProcessedDocument:
    text, encoding = _read_text_with_fallback(path)
    text = _clean_text(text)
    pages = [DocumentPage(page_number=1, text=text, char_count=len(text))] if text else []
    status = DocumentStatus.EMPTY_TEXT if not text else DocumentStatus.EXTRACTED

    document = _build_document(
        path=path,
        file_hash=file_hash,
        document_id=document_id,
        status=status,
        pages=pages,
        page_count=len(pages),
        char_count=len(text),
    )
    document.metadata["encoding"] = encoding
    return document


def _load_docx(path: Path, file_hash: str, document_id: str) -> ProcessedDocument:
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text = _clean_text(text)
    pages = [DocumentPage(page_number=1, text=text, char_count=len(text))] if text else []
    status = DocumentStatus.EMPTY_TEXT if not text else DocumentStatus.EXTRACTED

    return _build_document(
        path=path,
        file_hash=file_hash,
        document_id=document_id,
        status=status,
        pages=pages,
        page_count=len(pages),
        char_count=len(text),
    )


def _read_text_with_fallback(path: Path) -> tuple[str, str]:
    last_error: UnicodeDecodeError | None = None

    for encoding in TEXT_ENCODINGS:
        try:
            return path.read_text(encoding=encoding), encoding
        except UnicodeDecodeError as exc:
            last_error = exc

    if last_error is not None:
        raise last_error

    return "", TEXT_ENCODINGS[0]


def _build_document(
    *,
    path: Path,
    file_hash: str,
    document_id: str,
    status: DocumentStatus,
    requires_ocr: bool = False,
    pages: list[DocumentPage] | None = None,
    page_count: int = 0,
    char_count: int = 0,
    error_message: str | None = None,
) -> ProcessedDocument:
    resolved_path = path.resolve()

    return ProcessedDocument(
        document_id=document_id,
        file_name=path.name,
        file_path=str(resolved_path),
        file_type=path.suffix.lower().lstrip("."),
        file_hash=file_hash,
        extraction_status=status,
        requires_ocr=requires_ocr,
        page_count=page_count,
        char_count=char_count,
        processed_at=datetime.now(timezone.utc),
        pages=pages or [],
        metadata={"source": _source_label(resolved_path)},
        error_message=error_message,
    )


def _source_label(path: Path) -> str:
    try:
        return path.relative_to(PROJECT_ROOT).as_posix()
    except ValueError:
        return str(path)


def _clean_text(text: str) -> str:
    return text.replace("\x00", "").strip()


def _safe_error(exc: Exception) -> str:
    message = f"{type(exc).__name__}: {exc}"
    return message[:MAX_ERROR_MESSAGE_LENGTH]
